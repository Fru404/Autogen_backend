import io
import json
import os
import shutil
import tempfile
import threading
import uuid
import zipfile
from pathlib import Path

import requests

from fastapi import (
    BackgroundTasks,
    FastAPI,
    File,
    Form,
    HTTPException,
    UploadFile,
)

from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import Response

from docx import Document

from document_processor import (
    get_placeholders,
    read_excel,
    generate_document,
)


# ============================================================
# CONFIGURATION
# ============================================================

# Master switch.
#
# True  = UI is allowed to request PDFs
# False = PDF generation disabled
#
GENERATE_PDF = True


# ConvertAPI token.
#
# IMPORTANT:
# Put this in Render Environment Variables.
#
# DO NOT put the token in Next.js.
#
CONVERTAPI_TOKEN = os.getenv(
    "CONVERTAPI_TOKEN"
)


# ============================================================
# APP
# ============================================================

app = FastAPI(
    title="Autogen Document Generator",
    version="1.0.3"
)


# ============================================================
# CORS
# ============================================================

app.add_middleware(
    CORSMiddleware,

    allow_origins=[
        "*"
    ],

    allow_credentials=True,

    allow_methods=[
        "*"
    ],

    allow_headers=[
        "*"
    ],
)


# ============================================================
# JOB STORAGE
# ============================================================

jobs = {}

jobs_lock = threading.Lock()


def create_job(total):

    job_id = str(
        uuid.uuid4()
    )

    job = {
        "id": job_id,

        "status": "starting",

        "progress": 0,

        "completed": 0,

        "total": total,

        "current": "",

        "message": "Preparing...",

        "error": None,

        "zip": None,
    }

    with jobs_lock:

        jobs[job_id] = job

    return job_id


def update_job(
    job_id,
    **values
):

    with jobs_lock:

        if job_id in jobs:

            jobs[job_id].update(
                values
            )


def get_job(job_id):

    with jobs_lock:

        return jobs.get(
            job_id
        )


# ============================================================
# HEALTH
# ============================================================

@app.get("/health")
def health():

    return {
        "status": "ok",

        "service":
            "autogen-document-generator",

        "pdf_enabled":
            GENERATE_PDF,
    }


# ============================================================
# INSPECT TEMPLATE + EXCEL
# ============================================================

@app.post("/inspect")
async def inspect_files(
    template: UploadFile = File(...),
    excel: UploadFile = File(...)
):

    template_name = (
        template.filename or ""
    ).lower()

    excel_name = (
        excel.filename or ""
    ).lower()


    # --------------------------------------------------------
    # Validate template
    # --------------------------------------------------------

    if not template_name.endswith(
        ".docx"
    ):

        raise HTTPException(
            status_code=400,

            detail=
                "Template must be a .docx file."
        )


    # --------------------------------------------------------
    # Validate Excel
    # --------------------------------------------------------

    if not excel_name.endswith(
        (
            ".xlsx",
            ".xls"
        )
    ):

        raise HTTPException(
            status_code=400,

            detail=
                "Excel file must be .xlsx or .xls."
        )


    temp_dir = tempfile.mkdtemp()


    try:

        template_path = (
            Path(temp_dir)
            / "template.docx"
        )

        excel_path = (
            Path(temp_dir)
            / "clients.xlsx"
        )


        # ----------------------------------------------------
        # Save template
        # ----------------------------------------------------

        with open(
            template_path,
            "wb"
        ) as file:

            shutil.copyfileobj(
                template.file,
                file
            )


        # ----------------------------------------------------
        # Save Excel
        # ----------------------------------------------------

        with open(
            excel_path,
            "wb"
        ) as file:

            shutil.copyfileobj(
                excel.file,
                file
            )


        # ----------------------------------------------------
        # Read template
        # ----------------------------------------------------

        document = Document(
            str(template_path)
        )


        placeholders = (
            get_placeholders(
                document
            )
        )


        # ----------------------------------------------------
        # Read Excel
        # ----------------------------------------------------

        columns, records = read_excel(
            excel_path
        )


        # ----------------------------------------------------
        # Missing fields
        # ----------------------------------------------------

        missing_fields = [

            field

            for field in placeholders

            if field not in columns
        ]


        # ----------------------------------------------------
        # Extra fields
        # ----------------------------------------------------

        extra_columns = [

            column

            for column in columns

            if column not in placeholders
        ]


        return {

            "template": {

                "filename":
                    template.filename,

                "fields":
                    sorted(
                        placeholders
                    ),
            },

            "excel": {

                "filename":
                    excel.filename,

                "columns":
                    columns,

                "row_count":
                    len(records),

                "rows":
                    records,
            },

            "validation": {

                "missing_fields":
                    missing_fields,

                "extra_columns":
                    extra_columns,

                "ready":
                    len(missing_fields) == 0,
            },

            "settings": {

                "pdf_enabled":
                    GENERATE_PDF,
            },
        }


    except Exception as error:

        raise HTTPException(
            status_code=400,
            detail=str(error)
        )


    finally:

        shutil.rmtree(
            temp_dir,
            ignore_errors=True
        )


# ============================================================
# CONVERT DOCX → PDF
# ============================================================

def convert_docx_to_pdf(
    docx_path,
    output_folder
):

    if not CONVERTAPI_TOKEN:

        raise RuntimeError(
            "CONVERTAPI_TOKEN is not configured "
            "on the server."
        )


    docx_path = Path(
        docx_path
    )

    output_folder = Path(
        output_folder
    )


    output_folder.mkdir(
        parents=True,
        exist_ok=True
    )


    url = (
        "https://v2.convertapi.com/"
        "convert/docx/to/pdf"
    )


    with open(
        docx_path,
        "rb"
    ) as file:

        response = requests.post(

            url,

            headers={
                "Authorization":
                    f"Bearer {CONVERTAPI_TOKEN}"
            },

            files={
                "File": (
                    docx_path.name,
                    file,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                )
            },

            data={
                "StoreFile":
                    "false"
            },

            timeout=120,
        )


    if not response.ok:

        raise RuntimeError(
            "ConvertAPI error "
            f"{response.status_code}: "
            f"{response.text}"
        )


    result = response.json()


    files = result.get(
        "Files",
        []
    )


    if not files:

        raise RuntimeError(
            "ConvertAPI did not return "
            "a converted PDF."
        )


    converted = files[0]


    # --------------------------------------------------------
    # ConvertAPI with StoreFile=false returns Base64 content
    # --------------------------------------------------------

    file_data = converted.get(
        "FileData"
    )


    if file_data:

        import base64

        pdf_bytes = base64.b64decode(
            file_data
        )


    else:

        # ----------------------------------------------------
        # Fallback if ConvertAPI returns a URL
        # ----------------------------------------------------

        file_url = converted.get(
            "Url"
        )


        if not file_url:

            raise RuntimeError(
                "ConvertAPI returned neither "
                "PDF data nor a download URL."
            )


        pdf_response = requests.get(
            file_url,
            timeout=120
        )


        pdf_response.raise_for_status()


        pdf_bytes = (
            pdf_response.content
        )


    pdf_path = (

        output_folder

        /

        f"{docx_path.stem}.pdf"
    )


    pdf_path.write_bytes(
        pdf_bytes
    )


    return pdf_path

# ============================================================
# ACTUAL GENERATION
# ============================================================

def run_generation(
    job_id,
    template_bytes,
    records,
    generate_pdf
):

    total = len(records)


    try:

        update_job(
            job_id,

            status="generating",

            progress=0,

            completed=0,

            total=total,

            message=
                "Starting receipt generation..."
        )


        # ----------------------------------------------------
        # Temporary workspace
        # ----------------------------------------------------

        with tempfile.TemporaryDirectory() as temp:

            temp_dir = Path(temp)


            template_path = (
                temp_dir
                / "template.docx"
            )


            output_dir = (
                temp_dir
                / "receipts"
            )


            output_dir.mkdir(
                parents=True,
                exist_ok=True
            )


            template_path.write_bytes(
                template_bytes
            )


            generated_files = []


            # =================================================
            # RECEIPTS
            # =================================================

            for index, record in enumerate(
                records,
                start=1
            ):

                if not isinstance(
                    record,
                    dict
                ):

                    raise RuntimeError(
                        f"Row {index} is invalid."
                    )


                # ------------------------------------------------
                # Determine client name
                # ------------------------------------------------

                name = str(
                    record.get(
                        "name",
                        f"Client {index}"
                    )
                    or f"Client {index}"
                )


                # ------------------------------------------------
                # Tell UI what is happening
                # ------------------------------------------------

                update_job(

                    job_id,

                    status=
                        "generating",

                    progress=
                        int(
                            (
                                (index - 1)
                                / total
                            )
                            * 100
                        ),

                    completed=
                        index - 1,

                    total=
                        total,

                    current=
                        name,

                    message=
                        f"Creating receipt "
                        f"{index} of {total}"
                )


                # ------------------------------------------------
                # Generate DOCX
                # ------------------------------------------------

                docx_file = (
                    generate_document(

                        template_file=
                            str(
                                template_path
                            ),

                        output_folder=
                            str(
                                output_dir
                            ),

                        data=
                            record,

                        row_number=
                            index
                    )
                )


                generated_files.append(
                    docx_file
                )


                # ------------------------------------------------
                # PDF
                # ------------------------------------------------

                if generate_pdf:

                    update_job(

                        job_id,

                        current=
                            name,

                        message=
                            f"Creating PDF "
                            f"{index} of {total}"
                    )


                    pdf_file = (
                        convert_docx_to_pdf(

                            docx_file,

                            output_dir
                        )
                    )


                    generated_files.append(
                        pdf_file
                    )


                # ------------------------------------------------
                # Completed receipt
                # ------------------------------------------------

                update_job(

                    job_id,

                    progress=
                        int(
                            (
                                index
                                / total
                            )
                            * 100
                        ),

                    completed=
                        index,

                    total=
                        total,

                    current=
                        name,

                    message=
                        f"Receipt {index} "
                        f"of {total} completed"
                )


            # =================================================
            # ZIP
            # =================================================

            update_job(

                job_id,

                status=
                    "zipping",

                progress=
                    99,

                current="",

                message=
                    "Creating ZIP file..."
            )


            zip_buffer = io.BytesIO()


            with zipfile.ZipFile(

                zip_buffer,

                "w",

                compression=
                    zipfile.ZIP_DEFLATED

            ) as archive:


                for file_path in generated_files:

                    file_path = Path(
                        file_path
                    )


                    archive.write(

                        file_path,

                        arcname=
                            file_path.name
                    )


            zip_bytes = (
                zip_buffer.getvalue()
            )


            # =================================================
            # COMPLETE
            # =================================================

            update_job(

                job_id,

                status=
                    "completed",

                progress=
                    100,

                completed=
                    total,

                total=
                    total,

                current="",

                message=
                    f"All {total} receipts completed.",

                zip=
                    zip_bytes
            )


    except Exception as error:

        update_job(

            job_id,

            status=
                "error",

            error=
                str(error),

            message=
                "Generation failed."
        )


# ============================================================
# START GENERATION
# ============================================================

@app.post("/generate")
async def generate_documents(

    background_tasks:
        BackgroundTasks,

    template:
        UploadFile = File(...),

    data_json:
        str = Form(...),

    generate_pdf:
        bool = Form(False),
):


    # --------------------------------------------------------
    # Validate template
    # --------------------------------------------------------

    template_name = (
        template.filename or ""
    ).lower()


    if not template_name.endswith(
        ".docx"
    ):

        raise HTTPException(

            status_code=400,

            detail=
                "Template must be a .docx file."
        )


    # --------------------------------------------------------
    # Parse rows
    # --------------------------------------------------------

    try:

        records = json.loads(
            data_json
        )

    except json.JSONDecodeError:

        raise HTTPException(

            status_code=400,

            detail=
                "Invalid client data."
        )


    if not isinstance(
        records,
        list
    ):

        raise HTTPException(

            status_code=400,

            detail=
                "Client data must be a list."
        )


    if len(records) == 0:

        raise HTTPException(

            status_code=400,

            detail=
                "No clients were supplied."
        )


    # --------------------------------------------------------
    # PDF permission
    # --------------------------------------------------------

    if generate_pdf:

        if not GENERATE_PDF:

            raise HTTPException(

                status_code=400,

                detail=
                    "PDF generation is disabled."
            )


        if not CONVERTAPI_TOKEN:

            raise HTTPException(

                status_code=500,

                detail=
                    "PDF generation is enabled "
                    "but CONVERTAPI_TOKEN is missing."
            )


    # --------------------------------------------------------
    # Read template into memory
    # --------------------------------------------------------

    template_bytes = (
        await template.read()
    )


    # --------------------------------------------------------
    # Create job
    # --------------------------------------------------------

    job_id = create_job(
        len(records)
    )


    # --------------------------------------------------------
    # Start background work
    # --------------------------------------------------------

    background_tasks.add_task(

        run_generation,

        job_id,

        template_bytes,

        records,

        generate_pdf
    )


    # --------------------------------------------------------
    # Immediately respond
    # --------------------------------------------------------

    return {

        "job_id":
            job_id,

        "status":
            "starting",

        "total":
            len(records),
    }


# ============================================================
# PROGRESS
# ============================================================

@app.get(
    "/generate/{job_id}/progress"
)
def generation_progress(
    job_id: str
):

    job = get_job(
        job_id
    )


    if not job:

        raise HTTPException(

            status_code=404,

            detail=
                "Generation job not found."
        )


    return {

        "job_id":
            job["id"],

        "status":
            job["status"],

        "progress":
            job["progress"],

        "completed":
            job["completed"],

        "total":
            job["total"],

        "current":
            job["current"],

        "message":
            job["message"],

        "error":
            job["error"],

        "zip_ready":
            job["zip"] is not None,
    }


# ============================================================
# DOWNLOAD ZIP
# ============================================================

@app.get(
    "/generate/{job_id}/download"
)
def download_generation(
    job_id: str
):

    job = get_job(
        job_id
    )


    if not job:

        raise HTTPException(

            status_code=404,

            detail=
                "Generation job not found."
        )


    if job["status"] != "completed":

        raise HTTPException(

            status_code=400,

            detail=
                "Generation is not complete."
        )


    zip_bytes = job["zip"]


    if not zip_bytes:

        raise HTTPException(

            status_code=500,

            detail=
                "ZIP file is unavailable."
        )


    return Response(

        content=
            zip_bytes,

        media_type=
            "application/zip",

        headers={

            "Content-Disposition":
                'attachment; filename="Autogen_Receipts.zip"'
        }
    )